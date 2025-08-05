import streamlit as st
import pandas as pd
import tempfile
from fpdf import FPDF
import openai
import os

# Load your OpenAI key (you can set this securely in Streamlit Cloud or locally)
openai.api_key = st.secrets["OPENAI_API_KEY"] if "OPENAI_API_KEY" in st.secrets else os.getenv("OPENAI_API_KEY")

st.set_page_config(page_title="Estimator AI", layout="wide")
st.title("AI-Powered Job Estimator")

# --- Upload Supplier Pricing File ---
st.sidebar.header("Upload Your Supplier Pricing")
price_file = st.sidebar.file_uploader("Upload CSV (Material,Price)", type=["csv"])
custom_prices = {}
if price_file:
    try:
        df_prices = pd.read_csv(price_file)
        for _, row in df_prices.iterrows():
            custom_prices[row["Material"].lower()] = float(row["Price"])
    except Exception as e:
        st.sidebar.error("Failed to read file. Format must be: Material,Price")

# --- Upload Scope Document ---
st.sidebar.header("Upload Scope Document (Optional)")
scope_doc = st.sidebar.file_uploader("Upload .txt or .docx", type=["txt", "docx"])
if scope_doc is not None:
    if scope_doc.name.endswith(".txt"):
        job_description = scope_doc.read().decode("utf-8")
    else:
        import docx
        doc = docx.Document(scope_doc)
        job_description = "\n".join([p.text for p in doc.paragraphs])
    st.success("Loaded description from uploaded document")
else:
    job_description = st.text_area("Or type the job description here:", "")

# --- GPT Smart Estimator ---
def fetch_estimate_from_gpt(description):
    prompt = f"""
    You're an experienced construction estimator. Based on the job description below, provide estimated quantities for the following:
    - 2x4s (units)
    - Drywall Sheets (4x8 ft)
    - Nails (lbs)
    - Paint (gallons)
    - Concrete Bags (80lb)
    - Labor Hours (total)

    Job Description:
    """
    try:
        response = openai.ChatCompletion.create(
            model="gpt-4",
            messages=[{"role": "user", "content": prompt + description}]
        )
        return response.choices[0].message.content
    except Exception as e:
        st.error("GPT error: " + str(e))
        return None

if st.button("Generate Estimate with AI") and job_description:
   for line in result.split("\n"):

    if result:
        st.text_area("AI-Generated Estimate", result, height=200)

# --- Parse AI Output and Match with Pricing ---
        parsed_items = {}
        for line in result.split("
"):
            parts = line.split(":")
            if len(parts) == 2:
                material = parts[0].strip().lower()
                try:
                    quantity = float(parts[1].strip().split()[0])
                    parsed_items[material] = quantity
                except:
                    continue

        st.subheader("Estimate Summary")
        total_cost = 0.0
        summary_rows = []

        for material, qty in parsed_items.items():
            price = custom_prices.get(material, None)
            if price is not None:
                cost = qty * price
                total_cost += cost
                summary_rows.append((material.title(), qty, price, cost))
            else:
                summary_rows.append((material.title(), qty, "N/A", "N/A"))

        if summary_rows:
            df_summary = pd.DataFrame(summary_rows, columns=["Material", "Quantity", "Unit Price", "Total Cost"])
            st.table(df_summary)
            st.markdown(f"### Total Estimated Cost: **${total_cost:,.2f}**")

            # Optional: PDF export
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("Arial", size=12)
            pdf.cell(200, 10, txt="Estimate Summary", ln=True, align="C")
            pdf.ln(10)
            for mat, qty, unit, cost in summary_rows:
                pdf.cell(200, 10, txt=f"{mat}: Qty={qty}, Unit=${unit}, Total=${cost}", ln=True)
            pdf.ln(10)
            pdf.set_font("Arial", 'B', 14)
            pdf.cell(200, 10, txt=f"Total Estimate: ${total_cost:.2f}", ln=True)

            with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_file:
                pdf.output(tmp_file.name)
                with open(tmp_file.name, "rb") as file:
                    st.download_button(
                        label="Download Estimate as PDF",
                        data=file,
                        file_name="estimate_summary.pdf",
                        mime="application/pdf"
                    )
        else:
            st.warning("No materials matched or parsed from the GPT output.")
